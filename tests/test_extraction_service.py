from __future__ import annotations

from io import BytesIO

import fitz
import pytest
from docx import Document

from backend.app.modules.docx_extractor import DOCXExtractor, DOCXExtractionError
from backend.app.modules.extraction_service import ExtractionError, ExtractionService
from backend.app.modules.pdf_extractor import PDFExtractor, PDFExtractionError


def build_pdf_bytes() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "John Doe\nSenior Python Engineer\nFastAPI   Streamlit")
    pdf_bytes = document.tobytes()
    document.close()
    return pdf_bytes


def build_docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("John   Doe")
    document.add_paragraph("Senior Python Engineer")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "FastAPI"
    table.cell(0, 1).text = "Streamlit"
    table.cell(1, 0).text = "Python"
    table.cell(1, 1).text = "SQL"

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_extractor_extracts_text_and_preserves_line_breaks() -> None:
    extractor = PDFExtractor()

    text = extractor.extract(build_pdf_bytes())

    assert "John Doe" in text
    assert "Senior Python Engineer" in text


def test_docx_extractor_extracts_paragraphs_and_tables() -> None:
    extractor = DOCXExtractor()

    text = extractor.extract(build_docx_bytes())

    assert text.splitlines()[0] == "John Doe"
    assert "Senior Python Engineer" in text
    assert "FastAPI | Streamlit" in text


def test_extraction_service_normalizes_whitespace() -> None:
    service = ExtractionService()

    result = service.extract_text("John   Doe\n\n  Senior   Python   Engineer  ")

    assert result == "John Doe\n\nSenior Python Engineer"


def test_extraction_service_extracts_pdf_resume() -> None:
    service = ExtractionService()

    result = service.extract_from_file("resume.pdf", build_pdf_bytes())

    assert result.file_type == "pdf"
    assert result.filename == "resume.pdf"
    assert "John Doe" in result.text


def test_extraction_service_extracts_docx_resume() -> None:
    service = ExtractionService()

    result = service.extract_from_file("resume.docx", build_docx_bytes())

    assert result.file_type == "docx"
    assert result.filename == "resume.docx"
    assert "FastAPI | Streamlit" in result.text


@pytest.mark.parametrize("filename,content", [("resume.pdf", b"not-a-pdf"), ("resume.docx", b"not-a-docx")])
def test_extraction_service_handles_corrupted_files(filename: str, content: bytes) -> None:
    service = ExtractionService()

    with pytest.raises(ExtractionError):
        service.extract_from_file(filename, content)
