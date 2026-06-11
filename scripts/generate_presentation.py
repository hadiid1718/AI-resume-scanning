"""Generate project overview PPTX."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUTPUT = Path(__file__).resolve().parents[1] / "docs" / "AI_Resume_Scanning_Presentation.pptx"

NAVY = RGBColor(15, 23, 42)
INDIGO = RGBColor(99, 102, 241)
SLATE = RGBColor(100, 116, 139)
WHITE = RGBColor(255, 255, 255)


def set_title(slide, text: str, subtitle: str = "") -> None:
    title = slide.shapes.title
    title.text = text
    title.text_frame.paragraphs[0].font.size = Pt(32)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = NAVY
    if subtitle and len(slide.placeholders) > 1:
        sub = slide.placeholders[1]
        sub.text = subtitle
        sub.text_frame.paragraphs[0].font.size = Pt(16)
        sub.text_frame.paragraphs[0].font.color.rgb = SLATE


def add_bullets(slide, items: list[str], left=0.7, top=1.6, width=8.5, height=5.0) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(18)
        p.font.color.rgb = NAVY
        p.space_after = Pt(10)


def add_flow_box(slide, text: str, left, top, width=1.35, height=0.55, fill=INDIGO) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = NAVY
    tf = shape.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = 1


def add_arrow(slide, x1, y1, x2, y2) -> None:
    slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2)).line.color.rgb = SLATE


def build() -> Path:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Slide 1 — Title
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    set_title(
        s1,
        "AI Resume Scanning System",
        "Automated recruitment screening with LangGraph + Gemini API",
    )

    # Slide 2 — Problem & Ideas
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(s2, "Problem & Project Idea")
    add_bullets(
        s2,
        [
            "Problem: HR teams receive hundreds of resumes per role. Manual screening is slow, inconsistent, and biased.",
            "Pain points: missed qualified candidates, uneven evaluation, no structured comparison, heavy time cost.",
            "Who it is for: HR recruiters, hiring managers, and talent-acquisition teams in tech-driven organizations.",
            "What the agent solves: batch resume intake, structured extraction, 0–100 AI scoring, ranking, and PDF reports.",
            "Why we chose this: high-volume hiring needs repeatable, explainable, AI-assisted decisions—not guesswork.",
        ],
    )

    # Slide 3 — Solution
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(s3, "Solution Overview")
    add_bullets(
        s3,
        [
            "Streamlit dashboard: upload resumes (batch), track processing, search/filter, compare candidates.",
            "FastAPI backend: REST APIs for analyze, batch analyze, rankings, and report export.",
            "Gemini API: intelligent resume understanding + evaluation (skills, experience, education, verdict).",
            "SQLite storage: persist candidates, scores, recommendations, and full analysis payloads.",
            "Outputs: ranked leaderboard, recommendation labels, individual & batch HR reports (PDF/Excel).",
        ],
    )

    # Slide 4 — Architecture
    s4 = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(s4, "Architecture & Workflow Type")
    add_bullets(
        s4,
        [
            "Workflow type: LangGraph StateGraph — linear pipeline with shared ResumePipelineState.",
            "Pattern: deterministic DAG (Directed Acyclic Graph) with explicit nodes and typed state passing.",
            "Frontend: Streamlit  |  Backend: FastAPI  |  AI: LangGraph + Gemini  |  DB: SQLite",
            "Why LangGraph: clear stage separation, easy to extend (new nodes), observable status, batch-safe.",
            "Why not a single monolithic script: each stage (parse, extract, match, evaluate) can evolve independently.",
        ],
    )

    # Slide 5 — Flow diagram
    s5 = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(s5, "LangGraph Workflow — Nodes, Edges & Decisions")

    nodes = [
        ("START", 0.4, 1.5, SLATE),
        ("parse_resume", 1.5, 1.5, INDIGO),
        ("extract_candidate", 3.2, 1.5, INDIGO),
        ("parse_job", 5.0, 1.5, INDIGO),
        ("match", 6.7, 1.5, INDIGO),
        ("evaluate", 1.5, 3.0, INDIGO),
        ("rank", 3.2, 3.0, INDIGO),
        ("report", 5.0, 3.0, INDIGO),
        ("END", 6.7, 3.0, SLATE),
    ]
    for label, x, y, color in nodes:
        add_flow_box(s5, label, x, y, fill=color)

    # Row 1 arrows
    for (x1, x2) in [(1.0, 1.5), (2.85, 3.2), (4.55, 5.0), (6.35, 6.7)]:
        add_arrow(s5, x1 + 0.95, 1.75, x2, 1.75)
    # Down from match to evaluate
    add_arrow(s5, 7.35, 2.05, 2.2, 3.0)
    # Row 2 arrows
    for (x1, x2) in [(2.85, 3.2), (4.55, 5.0), (6.35, 6.7)]:
        add_arrow(s5, x1 + 0.95, 3.25, x2, 3.25)

    decision_box = s5.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9.0), Inches(2.8))
    tf = decision_box.text_frame
    tf.word_wrap = True
    decisions = [
        "Decision 1 (extract_candidate): Gemini available? → enrich profile; else rule-based extraction only.",
        "Decision 2 (evaluate): Gemini API key set? → full AI score 0–100; else rule-based fallback scoring.",
        "Edges: linear edges between all core nodes; no cycles — each resume flows once through the graph.",
        "Batch mode: graph invoked per resume, then RankingEngine sorts all candidates by AI score.",
    ]
    for idx, line in enumerate(decisions):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(14)
        p.font.color.rgb = NAVY
        p.space_after = Pt(6)

    prs.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"Created: {path}")
