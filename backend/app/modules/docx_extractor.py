from __future__ import annotations

import logging
from io import BytesIO

logger = logging.getLogger(__name__)


class DOCXExtractionError(Exception):
    pass


class DOCXExtractor:
    def extract(self, docx_bytes: bytes) -> str:
        try:
            from docx import Document
        except ModuleNotFoundError as exc:
            raise DOCXExtractionError(
                "DOCX support requires python-docx. Run: pip install python-docx"
            ) from exc

        try:
            document = Document(BytesIO(docx_bytes))
        except Exception as exc:
            logger.exception("Failed to open DOCX resume")
            raise DOCXExtractionError("The DOCX resume is corrupted or unreadable.") from exc

        try:
            lines: list[str] = []

            for paragraph in document.paragraphs:
                text = self._clean_line(paragraph.text)
                if text:
                    lines.append(text)

            for table in document.tables:
                for row in table.rows:
                    row_text = " | ".join(self._clean_line(cell.text) for cell in row.cells)
                    row_text = self._clean_line(row_text)
                    if row_text:
                        lines.append(row_text)

            return "\n".join(lines)
        except Exception as exc:
            logger.exception("Failed to extract text from DOCX resume")
            raise DOCXExtractionError("The DOCX resume could not be processed.") from exc

    def _clean_line(self, value: str) -> str:
        return " ".join(value.split())
